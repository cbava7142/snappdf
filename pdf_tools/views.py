import os
import shutil
import tempfile
import logging
import zipfile
from pathlib import Path

from django.conf import settings
from django.http import FileResponse, JsonResponse
from django.utils.text import get_valid_filename
from django.views.decorators.http import require_POST
from django.views.decorators.csrf import csrf_exempt
from pathlib import Path
from django.utils.text import get_valid_filename

from pypdf import PdfReader, PdfWriter
from pypdf.errors import PyPdfError
from pdf2image import convert_from_path, pdfinfo_from_path
from pdf2image.exceptions import (
    PDFInfoNotInstalledError,
    PDFPageCountError,
    PDFPopplerTimeoutError,
    PDFSyntaxError,
)

logger = logging.getLogger(__name__)

# --- Settings for Merge PDF ---
MAX_FILES = getattr(settings, "PDF_MERGE_MAX_FILES", 20)
MAX_FILE_SIZE = getattr(settings, "PDF_MERGE_MAX_FILE_SIZE", 50 * 1024 * 1024)

# --- Settings for PDF to JPG ---
JPG_MAX_FILE_SIZE = getattr(settings, "PDF_TO_JPG_MAX_FILE_SIZE", 50 * 1024 * 1024)
JPG_MAX_PAGES = getattr(settings, "PDF_TO_JPG_MAX_PAGES", 100)
JPG_DEFAULT_DPI = getattr(settings, "PDF_TO_JPG_DEFAULT_DPI", 150)
JPG_MAX_DPI = getattr(settings, "PDF_TO_JPG_MAX_DPI", 300)
JPG_QUALITY = getattr(settings, "PDF_TO_JPG_JPEG_QUALITY", 85)
POPPLER_TIMEOUT = getattr(settings, "PDF_TO_JPG_TIMEOUT", 60)


class PDFMergeError(Exception):
    """Raised for user-facing validation/processing errors."""
    def __init__(self, message, status=400):
        super().__init__(message)
        self.message = message
        self.status = status


class SelfCleaningFileResponse(FileResponse):
    """
    FileResponse that deletes its temp directory once the response is closed.
    """
    def __init__(self, *args, cleanup_dir=None, **kwargs):
        self._cleanup_dir = cleanup_dir
        super().__init__(*args, **kwargs)

    def close(self):
        try:
            super().close()  # closes the open file handle
        finally:
            if self._cleanup_dir:
                shutil.rmtree(self._cleanup_dir, ignore_errors=True)
                self._cleanup_dir = None


def _save_upload(upload, dest_path):
    """Write an uploaded file to disk in chunks and validate it looks like a PDF."""
    if upload.size == 0:
        raise PDFMergeError(f"'{upload.name}' is empty.")
    if upload.size > MAX_FILE_SIZE:
        raise PDFMergeError(
            f"'{upload.name}' exceeds the {MAX_FILE_SIZE // (1024 * 1024)} MB limit.",
            status=413,
        )

    # Cheap sanity check on the header ("%PDF-") before writing anything
    header = upload.read(5)
    upload.seek(0)
    if header != b"%PDF-":
        raise PDFMergeError(f"'{upload.name}' is not a valid PDF file.")

    with open(dest_path, "wb") as out:
        for chunk in upload.chunks():
            out.write(chunk)


def _validate_and_save_single_pdf(upload, dest_path):
    """Validate a single uploaded file for PDF to JPG conversion."""
    if not upload.name.lower().endswith(".pdf"):
        raise PDFMergeError("Only .pdf files are accepted.", status=415)
    if upload.size == 0:
        raise PDFMergeError("The uploaded file is empty.")
    if upload.size > JPG_MAX_FILE_SIZE:
        raise PDFMergeError(
            f"File exceeds the {JPG_MAX_FILE_SIZE // (1024 * 1024)} MB limit.", status=413
        )

    header = upload.read(5)
    upload.seek(0)
    if header != b"%PDF-":
        raise PDFMergeError("The uploaded file is not a valid PDF.", status=415)

    with open(dest_path, "wb") as out:
        for chunk in upload.chunks():
            out.write(chunk)


def _parse_dpi(raw_value):
    """Optional 'dpi' form field, clamped to a safe range to protect server memory."""
    if not raw_value:
        return JPG_DEFAULT_DPI
    try:
        dpi = int(raw_value)
    except ValueError:
        raise PDFMergeError("'dpi' must be an integer.")
    if not 72 <= dpi <= JPG_MAX_DPI:
        raise PDFMergeError(f"'dpi' must be between 72 and {JPG_MAX_DPI}.")
    return dpi


# ==========================================
# TOOL 1: MERGE PDF
# ==========================================
@csrf_exempt
@require_POST
def merge_pdf(request):
    """
    POST /merge/
    """
    uploads = request.FILES.getlist("files")

    if not uploads:
        return JsonResponse({"error": "No files uploaded. Use the 'files' field."}, status=400)
    if len(uploads) < 2:
        return JsonResponse({"error": "Upload at least two PDF files to merge."}, status=400)
    if len(uploads) > MAX_FILES:
        return JsonResponse({"error": f"Maximum {MAX_FILES} files allowed."}, status=400)

    work_dir = tempfile.mkdtemp(prefix="merge_", dir=getattr(settings, "PDF_TEMP_DIR", None))
    response = None

    try:
        saved_paths = []
        for index, upload in enumerate(uploads):
            path = os.path.join(work_dir, f"{index:03d}.pdf")
            _save_upload(upload, path)
            saved_paths.append((upload.name, path))

        writer = PdfWriter()
        try:
            for original_name, path in saved_paths:
                try:
                    reader = PdfReader(path)
                    if reader.is_encrypted:
                        raise PDFMergeError(f"'{original_name}' is password protected. Unlock it first.")
                    if len(reader.pages) == 0:
                        raise PDFMergeError(f"'{original_name}' contains no pages.")
                    writer.append(reader)
                except PDFMergeError:
                    raise
                except (PyPdfError, ValueError, KeyError, OSError, RecursionError) as exc:
                    logger.warning("Corrupted PDF %s: %s", original_name, exc)
                    raise PDFMergeError(f"'{original_name}' is corrupted or unreadable.")

            merged_path = os.path.join(work_dir, "merged.pdf")
            with open(merged_path, "wb") as merged_file:
                writer.write(merged_file)
        finally:
            writer.close()

        response = SelfCleaningFileResponse(
            open(merged_path, "rb"),
            as_attachment=True,
            filename="merged.pdf",
            content_type="application/pdf",
            cleanup_dir=work_dir,
        )
        return response

    except PDFMergeError as exc:
        return JsonResponse({"error": exc.message}, status=exc.status)
    except Exception:
        logger.exception("Unexpected error while merging PDFs")
        return JsonResponse({"error": "Internal error while merging PDFs."}, status=500)
    finally:
        if response is None:
            shutil.rmtree(work_dir, ignore_errors=True)


# ==========================================
# TOOL 2: PDF TO JPG
# ==========================================
@csrf_exempt
@require_POST
def pdf_to_jpg(request):
    """
    POST /pdf-to-jpg/
    """
    upload = request.FILES.get("file")
    if upload is None:
        return JsonResponse({"error": "No file uploaded. Use the 'file' field."}, status=400)

    work_dir = tempfile.mkdtemp(prefix="pdf2jpg_", dir=getattr(settings, "PDF_TEMP_DIR", None))
    response = None

    try:
        dpi = _parse_dpi(request.POST.get("dpi"))

        pdf_path = os.path.join(work_dir, "input.pdf")
        _validate_and_save_single_pdf(upload, pdf_path)

        try:
            info = pdfinfo_from_path(pdf_path, timeout=POPPLER_TIMEOUT)
            page_count = int(info["Pages"])
        except PDFInfoNotInstalledError:
            logger.error("Poppler is not installed on this server.")
            raise PDFMergeError("Server is missing PDF conversion tools.", status=500)
        except PDFPopplerTimeoutError:
            raise PDFMergeError("PDF analysis timed out.", status=408)
        except (PDFPageCountError, PDFSyntaxError, KeyError, ValueError):
            raise PDFMergeError("The PDF is corrupted, unreadable, or password protected.")

        if page_count < 1:
            raise PDFMergeError("The PDF contains no pages.")
        if page_count > JPG_MAX_PAGES:
            raise PDFMergeError(f"PDF has {page_count} pages; the maximum is {JPG_MAX_PAGES}.", status=413)

        base_name = get_valid_filename(Path(upload.name).stem) or "document"
        zip_path = os.path.join(work_dir, "images.zip")
        page_tmp = os.path.join(work_dir, "page.jpg")

        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_STORED) as zf:
            for page_no in range(1, page_count + 1):
                try:
                    images = convert_from_path(
                        pdf_path, dpi=dpi, first_page=page_no, last_page=page_no, timeout=POPPLER_TIMEOUT,
                    )
                except PDFPopplerTimeoutError:
                    raise PDFMergeError(f"Page {page_no} took too long to render.", status=408)
                except (PDFPageCountError, PDFSyntaxError):
                    raise PDFMergeError(f"Page {page_no} is corrupted and cannot be rendered.")

                if not images:
                    raise PDFMergeError(f"Page {page_no} could not be rendered.")

                image = images[0].convert("RGB")
                try:
                    image.save(page_tmp, "JPEG", quality=JPG_QUALITY, optimize=True)
                finally:
                    image.close()
                    for img in images:
                        img.close()

                zf.write(page_tmp, arcname=f"{base_name}_page_{page_no:03d}.jpg")
                os.remove(page_tmp)

        response = SelfCleaningFileResponse(
            open(zip_path, "rb"),
            as_attachment=True,
            filename=f"{base_name}_images.zip",
            content_type="application/zip",
            cleanup_dir=work_dir,
        )
        return response

    except PDFMergeError as exc:
        return JsonResponse({"error": exc.message}, status=exc.status)
    except Exception:
        logger.exception("Unexpected error during PDF to JPG conversion")
        return JsonResponse({"error": "Internal error while converting the PDF."}, status=500)
    finally:
        if response is None:
            shutil.rmtree(work_dir, ignore_errors=True)


# ======================================================================
# Tools 3 & 4: Protect PDF / Compress PDF
# ======================================================================
import subprocess

PROTECT_MAX_FILE_SIZE = getattr(settings, "PDF_PROTECT_MAX_FILE_SIZE", 50 * 1024 * 1024)
PROTECT_MAX_PASSWORD_LEN = getattr(settings, "PDF_PROTECT_MAX_PASSWORD_LENGTH", 128)
COMPRESS_MAX_FILE_SIZE = getattr(settings, "PDF_COMPRESS_MAX_FILE_SIZE", 100 * 1024 * 1024)
COMPRESS_TIMEOUT = getattr(settings, "PDF_COMPRESS_TIMEOUT", 120)
GS_BINARY = getattr(settings, "GHOSTSCRIPT_BINARY", "gs")

# Whitelist of user-selectable compression levels -> Ghostscript presets.
# Only these values ever reach the command line (no user-controlled arguments).
COMPRESS_LEVELS = {
    "low": "/printer",     # ~300 dpi, mild compression
    "medium": "/ebook",    # ~150 dpi, good balance (default)
    "high": "/screen",     # ~72 dpi, smallest file
}


def _save_pdf_upload(upload, dest_path, max_size):
    """Validate a single PDF upload (extension, size, magic bytes) and save it in chunks."""
    if not upload.name.lower().endswith(".pdf"):
        raise PDFMergeError("Only .pdf files are accepted.", status=415)
    if upload.size == 0:
        raise PDFMergeError("The uploaded file is empty.")
    if upload.size > max_size:
        raise PDFMergeError(
            f"File exceeds the {max_size // (1024 * 1024)} MB limit.", status=413
        )

    header = upload.read(5)
    upload.seek(0)
    if header != b"%PDF-":
        raise PDFMergeError("The uploaded file is not a valid PDF.", status=415)

    with open(dest_path, "wb") as out:
        for chunk in upload.chunks():
            out.write(chunk)


# ----------------------------------------------------------------------
# Tool 3: Protect PDF
# ----------------------------------------------------------------------
@csrf_exempt
@require_POST
def protect_pdf(request):
    """
    POST /api/protect-pdf/
    Form-data: file=<pdf>, password=<string>
    Returns: AES-256 encrypted PDF as an attachment, or JSON {"error": "..."}.
    """
    upload = request.FILES.get("file")
    password = request.POST.get("password", "")   # never log or echo this value

    # --- Input validation ---
    if upload is None:
        return JsonResponse({"error": "No file uploaded. Use the 'file' field."}, status=400)
    if not password:
        return JsonResponse({"error": "A 'password' is required."}, status=400)
    if len(password) > PROTECT_MAX_PASSWORD_LEN:
        return JsonResponse(
            {"error": f"Password must be at most {PROTECT_MAX_PASSWORD_LEN} characters."},
            status=400,
        )

    work_dir = tempfile.mkdtemp(
        prefix="protect_", dir=getattr(settings, "PDF_TEMP_DIR", None)
    )
    response = None

    try:
        # --- 1. Save upload under a generated name ---
        input_path = os.path.join(work_dir, "input.pdf")
        _save_pdf_upload(upload, input_path, PROTECT_MAX_FILE_SIZE)

        # --- 2. Read and validate the PDF ---
        try:
            reader = PdfReader(input_path)
            if reader.is_encrypted:
                raise PDFMergeError("This PDF is already password protected.")
            if len(reader.pages) == 0:
                raise PDFMergeError("The PDF contains no pages.")

            # clone_from copies pages, metadata and structure into the writer
            writer = PdfWriter(clone_from=reader)
        except PDFMergeError:
            raise
        except (PyPdfError, ValueError, KeyError, OSError, RecursionError) as exc:
            logger.warning("Unreadable PDF for protection: %s", exc)
            raise PDFMergeError("The PDF is corrupted or unreadable.")

        # --- 3. Encrypt (AES-256) and write the output ---
        output_path = os.path.join(work_dir, "protected.pdf")
        try:
            # owner password defaults to the user password when omitted
            writer.encrypt(user_password=password, algorithm="AES-256")
            with open(output_path, "wb") as out:
                writer.write(out)
        except ImportError:
            logger.error("'cryptography' package is missing; cannot use AES encryption.")
            raise PDFMergeError("Server encryption support is not configured.", status=500)
        except (PyPdfError, ValueError, OSError) as exc:
            logger.warning("Encryption failed: %s", exc)
            raise PDFMergeError("Failed to encrypt the PDF.", status=500)
        finally:
            writer.close()

        # --- 4. Serve + auto-delete the whole work_dir after the response ---
        base_name = get_valid_filename(Path(upload.name).stem) or "document"
        response = SelfCleaningFileResponse(
            open(output_path, "rb"),
            as_attachment=True,
            filename=f"{base_name}_protected.pdf",
            content_type="application/pdf",
            cleanup_dir=work_dir,
        )
        return response

    except PDFMergeError as exc:
        return JsonResponse({"error": exc.message}, status=exc.status)
    except Exception:
        logger.exception("Unexpected error while protecting PDF")
        return JsonResponse({"error": "Internal error while protecting the PDF."}, status=500)
    finally:
        if response is None:
            shutil.rmtree(work_dir, ignore_errors=True)


# ----------------------------------------------------------------------
# Tool 4: Compress PDF (Ghostscript)
# ----------------------------------------------------------------------
@csrf_exempt
@require_POST
def compress_pdf(request):
    """
    POST /api/compress-pdf/
    Form-data: file=<pdf>
               level=<low|medium|high> (optional, default "medium")
    Returns: compressed PDF as an attachment, or JSON {"error": "..."}.
    If compression would make the file larger, the original is returned unchanged.
    Response headers X-Original-Size / X-Compressed-Size report the byte sizes.
    """
    upload = request.FILES.get("file")
    if upload is None:
        return JsonResponse({"error": "No file uploaded. Use the 'file' field."}, status=400)

    level = request.POST.get("level", "medium").lower()
    if level not in COMPRESS_LEVELS:
        return JsonResponse(
            {"error": f"'level' must be one of: {', '.join(COMPRESS_LEVELS)}."}, status=400
        )

    work_dir = tempfile.mkdtemp(
        prefix="compress_", dir=getattr(settings, "PDF_TEMP_DIR", None)
    )
    response = None

    try:
        # --- 1. Save upload under a generated name ---
        input_path = os.path.join(work_dir, "input.pdf")
        output_path = os.path.join(work_dir, "output.pdf")
        _save_pdf_upload(upload, input_path, COMPRESS_MAX_FILE_SIZE)

        # Reject password-protected files early (Ghostscript can't open them).
        # If pypdf can't parse the file we let Ghostscript try to repair it.
        try:
            if PdfReader(input_path).is_encrypted:
                raise PDFMergeError("Password-protected PDFs cannot be compressed. Unlock it first.")
        except PDFMergeError:
            raise
        except Exception:
            pass

        # --- 2. Run Ghostscript (argument list, no shell => no injection risk) ---
        cmd = [
            GS_BINARY,
            "-sDEVICE=pdfwrite",
            "-dCompatibilityLevel=1.4",
            f"-dPDFSETTINGS={COMPRESS_LEVELS[level]}",
            "-dNOPAUSE",
            "-dBATCH",
            "-dQUIET",
            "-dSAFER",                       # block file-system access from PDF/PostScript
            f"-sOutputFile={output_path}",
            input_path,
        ]

        try:
            result = subprocess.run(
                cmd,
                cwd=work_dir,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=COMPRESS_TIMEOUT,    # child is killed on timeout
                check=False,
            )
        except FileNotFoundError:
            logger.error("Ghostscript binary '%s' not found.", GS_BINARY)
            raise PDFMergeError("Server is missing PDF compression tools.", status=500)
        except subprocess.TimeoutExpired:
            raise PDFMergeError("Compression timed out. Try a smaller file.", status=408)

        if result.returncode != 0:
            logger.warning(
                "Ghostscript failed (code %s): %s",
                result.returncode,
                result.stderr.decode("utf-8", errors="replace")[:500],
            )
            raise PDFMergeError("The PDF is corrupted or could not be compressed.")

        if not os.path.exists(output_path) or os.path.getsize(output_path) == 0:
            raise PDFMergeError("Compression produced no output. The PDF may be corrupted.")

        # --- 3. Never return a bigger file than we received ---
        original_size = os.path.getsize(input_path)
        compressed_size = os.path.getsize(output_path)
        served_path = output_path if compressed_size < original_size else input_path
        served_size = min(compressed_size, original_size)

        # --- 4. Serve + auto-delete the whole work_dir after the response ---
        base_name = get_valid_filename(Path(upload.name).stem) or "document"
        response = SelfCleaningFileResponse(
            open(served_path, "rb"),
            as_attachment=True,
            filename=f"{base_name}_compressed.pdf",
            content_type="application/pdf",
            cleanup_dir=work_dir,
        )
        response["X-Original-Size"] = str(original_size)
        response["X-Compressed-Size"] = str(served_size)
        return response

    except PDFMergeError as exc:
        return JsonResponse({"error": exc.message}, status=exc.status)
    except Exception:
        logger.exception("Unexpected error while compressing PDF")
        return JsonResponse({"error": "Internal error while compressing the PDF."}, status=500)
    finally:
        if response is None:
            shutil.rmtree(work_dir, ignore_errors=True)

# ======================================================================
# Batch 3: PDF to Text / PDF to Word / PDF to Speech
# ======================================================================

TEXT_MAX_FILE_SIZE = getattr(settings, "PDF_TEXT_MAX_FILE_SIZE", 50 * 1024 * 1024)
TEXT_MAX_PAGES = getattr(settings, "PDF_TEXT_MAX_PAGES", 500)
WORD_MAX_FILE_SIZE = getattr(settings, "PDF_WORD_MAX_FILE_SIZE", 30 * 1024 * 1024)
WORD_MAX_PAGES = getattr(settings, "PDF_WORD_MAX_PAGES", 100)
SPEECH_MAX_FILE_SIZE = getattr(settings, "PDF_SPEECH_MAX_FILE_SIZE", 20 * 1024 * 1024)
SPEECH_MAX_CHARS = getattr(settings, "PDF_SPEECH_MAX_CHARS", 20000)
SPEECH_TIMEOUT = getattr(settings, "PDF_SPEECH_TIMEOUT", 30)

NO_TEXT_MESSAGE = (
    "No extractable text found. The PDF is probably scanned images and needs OCR first."
)


# ----------------------------------------------------------------------
# Shared helpers
# ----------------------------------------------------------------------
def _run_single_pdf_tool(request, *, prefix, max_size, processor):
    """
    Common pipeline for single-file tools:
      validate 'file' -> save to private temp dir -> processor() -> SelfCleaningFileResponse.

    processor(request, input_path, work_dir) must return
    (output_path, file_extension, content_type) or raise PDFMergeError.
    The whole work_dir is deleted after the response closes, or immediately on error.
    """
    upload = request.FILES.get("file")
    if upload is None:
        return JsonResponse({"error": "No file uploaded. Use the 'file' field."}, status=400)

    work_dir = tempfile.mkdtemp(prefix=prefix, dir=getattr(settings, "PDF_TEMP_DIR", None))
    response = None

    try:
        input_path = os.path.join(work_dir, "input.pdf")   # generated name, never user input
        _save_pdf_upload(upload, input_path, max_size)

        output_path, extension, content_type = processor(request, input_path, work_dir)

        if not os.path.exists(output_path) or os.path.getsize(output_path) == 0:
            raise PDFMergeError("Conversion produced no output. The PDF may be corrupted.")

        base_name = get_valid_filename(Path(upload.name).stem) or "document"
        response = SelfCleaningFileResponse(
            open(output_path, "rb"),
            as_attachment=True,
            filename=f"{base_name}{extension}",
            content_type=content_type,
            cleanup_dir=work_dir,
        )
        return response

    except PDFMergeError as exc:
        return JsonResponse({"error": exc.message}, status=exc.status)
    except Exception:
        logger.exception("Unexpected error in %s tool", prefix.rstrip("_"))
        return JsonResponse({"error": "Internal error while processing the PDF."}, status=500)
    finally:
        if response is None:
            shutil.rmtree(work_dir, ignore_errors=True)


def _check_pdf_readable(pdf_path, max_pages):
    """Open with pypdf: reject encrypted, empty, oversized or corrupted PDFs. Returns the reader."""
    try:
        reader = PdfReader(pdf_path)
        if reader.is_encrypted:
            raise PDFMergeError("The PDF is password protected. Unlock it first.")
        page_count = len(reader.pages)
    except PDFMergeError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError, RecursionError) as exc:
        logger.warning("Unreadable PDF: %s", exc)
        raise PDFMergeError("The PDF is corrupted or unreadable.")

    if page_count == 0:
        raise PDFMergeError("The PDF contains no pages.")
    if page_count > max_pages:
        raise PDFMergeError(
            f"PDF has {page_count} pages; the maximum is {max_pages}.", status=413
        )
    return reader


def _extract_pdf_text(pdf_path, max_pages):
    """
    Extract text from every page with pypdf.
    Returns a list of page strings. Raises PDFMergeError (422) if no page has text
    (typical for scanned PDFs), or 400 if a page is corrupted.
    """
    reader = _check_pdf_readable(pdf_path, max_pages)

    pages = []
    for page_no, page in enumerate(reader.pages, start=1):
        try:
            pages.append((page.extract_text() or "").strip())
        except (PyPdfError, ValueError, KeyError, OSError, RecursionError) as exc:
            logger.warning("Text extraction failed on page %s: %s", page_no, exc)
            raise PDFMergeError(f"Page {page_no} is corrupted and its text cannot be read.")

    if not any(pages):
        raise PDFMergeError(NO_TEXT_MESSAGE, status=422)
    return pages


# ----------------------------------------------------------------------
# Tool 5: PDF to Text
# ----------------------------------------------------------------------
def _process_pdf_to_text(request, input_path, work_dir):
    pages = _extract_pdf_text(input_path, TEXT_MAX_PAGES)

    output_path = os.path.join(work_dir, "output.txt")
    with open(output_path, "w", encoding="utf-8") as out:
        out.write("\n\n".join(pages))   # blank line between pages
    return output_path, ".txt", "text/plain; charset=utf-8"


@csrf_exempt
@require_POST
def pdf_to_text(request):
    """
    POST /api/pdf-to-text/
    Form-data: file=<pdf>
    Returns: UTF-8 .txt attachment, or JSON {"error": "..."}.
    """
    return _run_single_pdf_tool(
        request, prefix="pdf2txt_", max_size=TEXT_MAX_FILE_SIZE, processor=_process_pdf_to_text
    )


# ----------------------------------------------------------------------
# Tool 6: PDF to Word (pdf2docx)
# ----------------------------------------------------------------------
def _process_pdf_to_word(request, input_path, work_dir):
    _check_pdf_readable(input_path, WORD_MAX_PAGES)   # encrypted / corrupt / too many pages

    try:
        from pdf2docx import Converter    # lazy import: heavy (PyMuPDF, OpenCV)
    except ImportError:
        logger.error("pdf2docx is not installed.")
        raise PDFMergeError("Server is missing PDF to Word conversion tools.", status=500)

    output_path = os.path.join(work_dir, "output.docx")
    converter = None
    try:
        converter = Converter(input_path)
        converter.convert(output_path, start=0, end=None)   # single process, all pages
    except Exception as exc:   # pdf2docx/PyMuPDF raise many different exception types
        logger.warning("pdf2docx conversion failed: %s", exc)
        raise PDFMergeError("The PDF could not be converted. It may be corrupted.")
    finally:
        if converter is not None:
            try:
                converter.close()
            except Exception:
                pass

    return (
        output_path,
        ".docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@csrf_exempt
@require_POST
def pdf_to_word(request):
    """
    POST /api/pdf-to-word/
    Form-data: file=<pdf>
    Returns: .docx attachment, or JSON {"error": "..."}.
    """
    return _run_single_pdf_tool(
        request, prefix="pdf2word_", max_size=WORD_MAX_FILE_SIZE, processor=_process_pdf_to_word
    )


# ----------------------------------------------------------------------
# Tool 7: PDF to Speech (pypdf + gTTS)
# ----------------------------------------------------------------------
def _process_pdf_to_speech(request, input_path, work_dir):
    try:
        from gtts import gTTS                 # lazy import
        from gtts.lang import tts_langs
        from gtts.tts import gTTSError
    except ImportError:
        logger.error("gTTS is not installed.")
        raise PDFMergeError("Server is missing text-to-speech tools.", status=500)

    # Validate the optional 'lang' field against gTTS's supported languages
    lang = request.POST.get("lang", "en").strip().lower()
    if lang not in tts_langs():
        raise PDFMergeError(f"Unsupported language '{lang}'.")

    pages = _extract_pdf_text(input_path, TEXT_MAX_PAGES)

    # Collapse whitespace/newlines so speech flows naturally
    text = " ".join(" ".join(pages).split())
    if not text:
        raise PDFMergeError(NO_TEXT_MESSAGE, status=422)
    if len(text) > SPEECH_MAX_CHARS:
        raise PDFMergeError(
            f"The PDF contains {len(text)} characters; the limit for audio is {SPEECH_MAX_CHARS}.",
            status=413,
        )

    output_path = os.path.join(work_dir, "output.mp3")
    try:
        gTTS(text=text, lang=lang, timeout=SPEECH_TIMEOUT).save(output_path)
    except AssertionError:
        # gTTS raises AssertionError("No text to speak") when nothing is speakable
        raise PDFMergeError(NO_TEXT_MESSAGE, status=422)
    except gTTSError as exc:
        logger.warning("gTTS request failed: %s", exc)
        raise PDFMergeError("The speech service is unavailable. Try again later.", status=502)
    except (ValueError, OSError) as exc:
        logger.warning("Speech generation failed: %s", exc)
        raise PDFMergeError("Failed to generate audio from this PDF.", status=500)

    return output_path, ".mp3", "audio/mpeg"


@csrf_exempt
@require_POST
def pdf_to_speech(request):
    """
    POST /api/pdf-to-speech/
    Form-data: file=<pdf>
               lang=<gTTS language code> (optional, default "en")
    Returns: .mp3 attachment, or JSON {"error": "..."}.
    """
    return _run_single_pdf_tool(
        request, prefix="pdf2speech_", max_size=SPEECH_MAX_FILE_SIZE, processor=_process_pdf_to_speech
    )

# ======================================================================
# Batch 4: Split PDF (extract pages) / Unlock PDF
# ======================================================================
import re

from pypdf.errors import DependencyError

SPLIT_MAX_FILE_SIZE = getattr(settings, "PDF_SPLIT_MAX_FILE_SIZE", 100 * 1024 * 1024)
SPLIT_MAX_PAGES = getattr(settings, "PDF_SPLIT_MAX_PAGES", 2000)
SPLIT_MAX_RANGES = getattr(settings, "PDF_SPLIT_MAX_RANGES", 200)
SPLIT_MAX_PAGES_STRING = getattr(settings, "PDF_SPLIT_MAX_PAGES_STRING", 1000)
UNLOCK_MAX_FILE_SIZE = getattr(settings, "PDF_UNLOCK_MAX_FILE_SIZE", 100 * 1024 * 1024)
UNLOCK_MAX_PASSWORD_LEN = getattr(settings, "PDF_UNLOCK_MAX_PASSWORD_LENGTH", 128)

# Matches "5" or "5-7" (digit count capped so huge numbers are rejected cheaply)
_PAGE_TOKEN_RE = re.compile(r"^(\d{1,6})(?:-(\d{1,6}))?$")


# ----------------------------------------------------------------------
# Page-range helpers
# ----------------------------------------------------------------------
def _parse_page_ranges(raw):
    """
    Parse a string like "1,3,5-7" into [(1, 1), (3, 3), (5, 7)].
    Validates syntax only; bounds are checked once the page count is known.
    Raises PDFMergeError (400) on invalid input.
    """
    raw = (raw or "").strip()
    if not raw:
        raise PDFMergeError("'pages' is required, for example: 1,3,5-7.")
    if len(raw) > SPLIT_MAX_PAGES_STRING:
        raise PDFMergeError("'pages' value is too long.")

    tokens = raw.replace(" ", "").split(",")
    if len(tokens) > SPLIT_MAX_RANGES:
        raise PDFMergeError(f"Too many page ranges; the maximum is {SPLIT_MAX_RANGES}.")

    ranges = []
    for token in tokens:
        match = _PAGE_TOKEN_RE.match(token)
        if not match:
            raise PDFMergeError(
                f"Invalid page selection '{token}'. Use numbers and ranges like 1,3,5-7."
            )
        start = int(match.group(1))
        end = int(match.group(2)) if match.group(2) else start
        if start < 1 or end < 1:
            raise PDFMergeError("Page numbers start at 1.")
        if end < start:
            raise PDFMergeError(f"Invalid range '{token}': the start must not exceed the end.")
        ranges.append((start, end))
    return ranges


def _expand_page_ranges(ranges, total_pages):
    """
    Validate ranges against the PDF length and expand them into unique,
    zero-based page indexes, keeping the order the user typed.
    Bounds are checked BEFORE expanding, so a range like 1-999999 can never
    allocate a huge list.
    """
    indexes, seen = [], set()
    for start, end in ranges:
        if end > total_pages:
            bad = end if start <= total_pages else start
            raise PDFMergeError(
                f"Page {bad} is out of range; the PDF has {total_pages} page(s)."
            )
        for page_no in range(start, end + 1):
            if page_no not in seen:
                seen.add(page_no)
                indexes.append(page_no - 1)
    return indexes


# ----------------------------------------------------------------------
# Tool 8: Split PDF / Extract Pages
# ----------------------------------------------------------------------
def _process_split_pdf(request, input_path, work_dir):
    ranges = _parse_page_ranges(request.POST.get("pages"))   # fail fast on bad syntax

    # Rejects encrypted, corrupted, empty and oversized PDFs
    reader = _check_pdf_readable(input_path, SPLIT_MAX_PAGES)
    indexes = _expand_page_ranges(ranges, len(reader.pages))

    output_path = os.path.join(work_dir, "output.pdf")
    writer = PdfWriter()
    try:
        for index in indexes:
            writer.add_page(reader.pages[index])
        with open(output_path, "wb") as out:
            writer.write(out)
    except (PyPdfError, ValueError, KeyError, IndexError, OSError, RecursionError) as exc:
        logger.warning("Page extraction failed: %s", exc)
        raise PDFMergeError("Some selected pages are corrupted and cannot be extracted.")
    finally:
        writer.close()

    return output_path, "_split.pdf", "application/pdf"


@csrf_exempt
@require_POST
def split_pdf(request):
    """
    POST /api/split-pdf/
    Form-data: file=<pdf>, pages="1,3,5-7" (1-indexed; order kept, duplicates removed)
    Returns: new PDF containing only the selected pages, or JSON {"error": "..."}.
    """
    return _run_single_pdf_tool(
        request, prefix="split_", max_size=SPLIT_MAX_FILE_SIZE, processor=_process_split_pdf
    )


# ----------------------------------------------------------------------
# Tool 9: Unlock PDF
# ----------------------------------------------------------------------
def _process_unlock_pdf(request, input_path, work_dir):
    # The field must be present; an empty value is allowed because some PDFs
    # only have an owner password (restrictions) and open with an empty user password.
    password = request.POST.get("password")   # never log or echo this value
    if password is None:
        raise PDFMergeError("A 'password' field is required.")
    if len(password) > UNLOCK_MAX_PASSWORD_LEN:
        raise PDFMergeError(
            f"Password must be at most {UNLOCK_MAX_PASSWORD_LEN} characters."
        )

    # --- Open the file and confirm it is actually encrypted ---
    try:
        reader = PdfReader(input_path)
        if not reader.is_encrypted:
            raise PDFMergeError("This PDF is not password protected.")

        # decrypt() returns PasswordType (0 = NOT_DECRYPTED, 1 = user, 2 = owner)
        if not reader.decrypt(password):
            raise PDFMergeError("Incorrect password.")

        if len(reader.pages) == 0:
            raise PDFMergeError("The PDF contains no pages.")
        writer = PdfWriter(clone_from=reader)   # decrypted copy, no encryption applied
    except PDFMergeError:
        raise
    except DependencyError:
        logger.error("'cryptography' package is missing; cannot decrypt AES PDFs.")
        raise PDFMergeError("Server decryption support is not configured.", status=500)
    except (PyPdfError, NotImplementedError, ValueError, KeyError, OSError, RecursionError) as exc:
        logger.warning("Unreadable or unsupported encrypted PDF: %s", exc)
        raise PDFMergeError("The PDF is corrupted or uses an unsupported encryption method.")

    # --- Write the unencrypted copy ---
    output_path = os.path.join(work_dir, "output.pdf")
    try:
        with open(output_path, "wb") as out:
            writer.write(out)
    except (PyPdfError, ValueError, KeyError, OSError, RecursionError) as exc:
        logger.warning("Writing unlocked PDF failed: %s", exc)
        raise PDFMergeError("The PDF is corrupted and cannot be unlocked.")
    finally:
        writer.close()

    return output_path, "_unlocked.pdf", "application/pdf"


@csrf_exempt
@require_POST
def unlock_pdf(request):
    """
    POST /api/unlock-pdf/
    Form-data: file=<encrypted pdf>, password=<string>
    Returns: decrypted PDF, or 400 JSON {"error": "..."} for a wrong password
    or a PDF that was never encrypted.
    """
    return _run_single_pdf_tool(
        request, prefix="unlock_", max_size=UNLOCK_MAX_FILE_SIZE, processor=_process_unlock_pdf
    )



# ======================================================================
# Batch 5: PDF to Excel / Excel to PDF / Image to PDF / Crop PDF
# ======================================================================
import math
from xml.sax.saxutils import escape as _xml_escape

from pypdf.generic import RectangleObject

EXCEL_MAX_FILE_SIZE = getattr(settings, "PDF_EXCEL_MAX_FILE_SIZE", 50 * 1024 * 1024)
EXCEL_MAX_PAGES = getattr(settings, "PDF_EXCEL_MAX_PAGES", 200)
EXCEL_MAX_CELLS = getattr(settings, "PDF_EXCEL_MAX_CELLS", 200_000)
XLSX_MAX_FILE_SIZE = getattr(settings, "EXCEL_PDF_MAX_FILE_SIZE", 20 * 1024 * 1024)
XLSX_MAX_SHEETS = getattr(settings, "EXCEL_PDF_MAX_SHEETS", 20)
XLSX_MAX_CELLS = getattr(settings, "EXCEL_PDF_MAX_CELLS", 50_000)
XLSX_MAX_CELL_CHARS = getattr(settings, "EXCEL_PDF_MAX_CELL_CHARS", 60)
PDF_FONT_PATH = getattr(settings, "PDF_FONT_PATH", None)
IMAGE_MAX_FILE_SIZE = getattr(settings, "IMAGE_PDF_MAX_FILE_SIZE", 30 * 1024 * 1024)
IMAGE_MAX_PIXELS = getattr(settings, "IMAGE_PDF_MAX_PIXELS", 50_000_000)
CROP_MAX_FILE_SIZE = getattr(settings, "PDF_CROP_MAX_FILE_SIZE", 100 * 1024 * 1024)
CROP_MAX_PAGES = getattr(settings, "PDF_CROP_MAX_PAGES", 2000)

XLSX_EXTENSIONS = {".xlsx", ".xlsm"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ----------------------------------------------------------------------
# Shared helper for NON-PDF inputs (mirrors _run_single_pdf_tool)
# ----------------------------------------------------------------------
def _is_zip_header(header):
    """.xlsx files are ZIP containers."""
    return header.startswith(b"PK\x03\x04")


def _is_image_header(header):
    return (
        header.startswith(b"\xff\xd8\xff")                          # JPEG
        or header.startswith(b"\x89PNG\r\n\x1a\n")                  # PNG
        or header.startswith(b"GIF8")                               # GIF
        or header.startswith(b"BM")                                 # BMP
        or header.startswith(b"II*\x00") or header.startswith(b"MM\x00*")   # TIFF
        or (header[:4] == b"RIFF" and header[8:12] == b"WEBP")      # WebP
    )


def _save_typed_upload(upload, work_dir, *, max_size, allowed_exts, magic_check, label):
    """
    Validate a single non-PDF upload (extension allowlist, size, magic bytes) and save it
    under a generated name that keeps only the validated extension. Returns the saved path.
    """
    ext = Path(upload.name).suffix.lower()
    if ext not in allowed_exts:
        raise PDFMergeError(
            f"Only {label} files are accepted ({', '.join(sorted(allowed_exts))}).", status=415
        )
    if upload.size == 0:
        raise PDFMergeError("The uploaded file is empty.")
    if upload.size > max_size:
        raise PDFMergeError(
            f"File exceeds the {max_size // (1024 * 1024)} MB limit.", status=413
        )

    header = upload.read(12)
    upload.seek(0)
    if not magic_check(header):
        raise PDFMergeError(f"The uploaded file is not a valid {label} file.", status=415)

    dest_path = os.path.join(work_dir, f"input{ext}")   # ext is from our allowlist, never raw user input
    with open(dest_path, "wb") as out:
        for chunk in upload.chunks():
            out.write(chunk)
    return dest_path


def _run_single_file_tool(request, *, prefix, max_size, allowed_exts, magic_check, label, processor):
    """
    Same pipeline as _run_single_pdf_tool, for tools whose INPUT is not a PDF.
    processor(request, input_path, work_dir) -> (output_path, extension, content_type).
    """
    upload = request.FILES.get("file")
    if upload is None:
        return JsonResponse({"error": "No file uploaded. Use the 'file' field."}, status=400)

    work_dir = tempfile.mkdtemp(prefix=prefix, dir=getattr(settings, "PDF_TEMP_DIR", None))
    response = None

    try:
        input_path = _save_typed_upload(
            upload, work_dir, max_size=max_size, allowed_exts=allowed_exts,
            magic_check=magic_check, label=label,
        )

        output_path, extension, content_type = processor(request, input_path, work_dir)

        if not os.path.exists(output_path) or os.path.getsize(output_path) == 0:
            raise PDFMergeError("Conversion produced no output. The file may be corrupted.")

        base_name = get_valid_filename(Path(upload.name).stem) or "document"
        response = SelfCleaningFileResponse(
            open(output_path, "rb"),
            as_attachment=True,
            filename=f"{base_name}{extension}",
            content_type=content_type,
            cleanup_dir=work_dir,
        )
        return response

    except PDFMergeError as exc:
        return JsonResponse({"error": exc.message}, status=exc.status)
    except Exception:
        logger.exception("Unexpected error in %s tool", prefix.rstrip("_"))
        return JsonResponse({"error": "Internal error while processing the file."}, status=500)
    finally:
        if response is None:
            shutil.rmtree(work_dir, ignore_errors=True)


# ----------------------------------------------------------------------
# Tool 10: PDF to Excel (pdfplumber + openpyxl)
# ----------------------------------------------------------------------
_NUMERIC_RE = re.compile(r"^-?\d+(\.\d+)?$")


def _excel_cell(value, illegal_re):
    """Normalise a pdfplumber cell for openpyxl: strip illegal chars, cast clean numbers."""
    if value is None:
        return None
    text = illegal_re.sub("", str(value)).strip()
    if not text:
        return None
    if _NUMERIC_RE.match(text):
        try:
            return int(text) if "." not in text else float(text)
        except ValueError:
            pass
    return text[:32767]   # Excel's per-cell character limit


def _process_pdf_to_excel(request, input_path, work_dir):
    _check_pdf_readable(input_path, EXCEL_MAX_PAGES)   # encrypted / corrupt / too many pages

    try:
        import pdfplumber                                    # lazy import
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    except ImportError:
        logger.error("pdfplumber/openpyxl not installed.")
        raise PDFMergeError("Server is missing PDF to Excel conversion tools.", status=500)

    workbook = Workbook()
    workbook.remove(workbook.active)   # start with no sheets
    cell_count = 0

    try:
        with pdfplumber.open(input_path) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables() or []
                wrote_table = False

                for table_no, table in enumerate(tables, start=1):
                    rows = [r for r in table if r and any(c not in (None, "") for c in r)]
                    if not rows:
                        continue
                    sheet = workbook.create_sheet(title=f"P{page_no}_T{table_no}"[:31])
                    for row in rows:
                        sheet.append([_excel_cell(c, ILLEGAL_CHARACTERS_RE) for c in row])
                        cell_count += len(row)
                        if cell_count > EXCEL_MAX_CELLS:
                            raise PDFMergeError(
                                f"The PDF contains more than {EXCEL_MAX_CELLS} cells.", status=413
                            )
                    wrote_table = True

                # No tables on this page: fall back to text, one line per row
                if not wrote_table:
                    text = (page.extract_text() or "").strip()
                    if text:
                        sheet = workbook.create_sheet(title=f"P{page_no}_Text"[:31])
                        for line in text.splitlines():
                            cell = _excel_cell(line, ILLEGAL_CHARACTERS_RE)
                            if cell is not None:
                                sheet.append([cell])
                                cell_count += 1
    except PDFMergeError:
        raise
    except Exception as exc:   # pdfplumber/pdfminer raise many exception types
        logger.warning("pdfplumber failed: %s", exc)
        raise PDFMergeError("The PDF is corrupted or its content cannot be read.")

    if not workbook.sheetnames:
        raise PDFMergeError(NO_TEXT_MESSAGE, status=422)

    output_path = os.path.join(work_dir, "output.xlsx")
    try:
        workbook.save(output_path)
    except (OSError, ValueError) as exc:
        logger.warning("Writing xlsx failed: %s", exc)
        raise PDFMergeError("Failed to write the Excel file.", status=500)

    return output_path, ".xlsx", EXCEL_MIME


@csrf_exempt
@require_POST
def pdf_to_excel(request):
    """
    POST /pdf-to-excel/
    Form-data: file=<pdf>
    Returns: .xlsx with one sheet per detected table (text fallback per page), or JSON error.
    """
    return _run_single_pdf_tool(
        request, prefix="pdf2xlsx_", max_size=EXCEL_MAX_FILE_SIZE, processor=_process_pdf_to_excel
    )


# ----------------------------------------------------------------------
# Tool 11: Excel to PDF (openpyxl + reportlab, pure Python)
# ----------------------------------------------------------------------
def _xlsx_cell_text(value):
    """Render an openpyxl cell value as a short display string."""
    if value is None:
        return ""
    if hasattr(value, "isoformat"):            # date / datetime / time
        text = value.isoformat(sep=" ") if hasattr(value, "hour") else value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = " ".join(text.split())              # collapse newlines/whitespace
    if len(text) > XLSX_MAX_CELL_CHARS:
        text = text[: XLSX_MAX_CELL_CHARS - 1] + "…"
    return text


def _register_pdf_font():
    """Use a TTF from settings.PDF_FONT_PATH if present (Unicode), else Helvetica."""
    if PDF_FONT_PATH and os.path.exists(PDF_FONT_PATH):
        try:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            if "SnapPDFSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("SnapPDFSans", PDF_FONT_PATH))
            return "SnapPDFSans", "SnapPDFSans"
        except Exception as exc:
            logger.warning("Could not register PDF font %s: %s", PDF_FONT_PATH, exc)
    return "Helvetica", "Helvetica-Bold"


def _process_excel_to_pdf(request, input_path, work_dir):
    try:
        from openpyxl import load_workbook                  # lazy imports
        from openpyxl.utils.exceptions import InvalidFileException
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
        )
    except ImportError:
        logger.error("openpyxl/reportlab not installed.")
        raise PDFMergeError("Server is missing Excel to PDF conversion tools.", status=500)

    # --- Read the workbook (cached formula values, streaming mode) ---
    try:
        workbook = load_workbook(input_path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        logger.warning("Unreadable xlsx: %s", exc)
        raise PDFMergeError("The Excel file is corrupted or unreadable.")
    except Exception as exc:
        logger.warning("Unexpected xlsx load failure: %s", exc)
        raise PDFMergeError("The Excel file is corrupted or unreadable.")

    if len(workbook.sheetnames) > XLSX_MAX_SHEETS:
        raise PDFMergeError(
            f"Workbook has {len(workbook.sheetnames)} sheets; the maximum is {XLSX_MAX_SHEETS}.",
            status=413,
        )

    sheets = []          # [(title, rows)]
    total_cells = 0
    try:
        for sheet in workbook.worksheets:
            rows, used_cols = [], 0
            for raw_row in sheet.iter_rows(values_only=True):
                row = [_xlsx_cell_text(v) for v in raw_row]
                last = max((i + 1 for i, v in enumerate(row) if v), default=0)
                rows.append(row)
                used_cols = max(used_cols, last)
                total_cells += len(row)
                if total_cells > XLSX_MAX_CELLS:
                    raise PDFMergeError(
                        f"Workbook has more than {XLSX_MAX_CELLS} cells; split it first.", status=413
                    )
            # Trim trailing empty rows and columns
            while rows and not any(rows[-1]):
                rows.pop()
            rows = [(r + [""] * used_cols)[:used_cols] for r in rows]
            if rows and used_cols:
                sheets.append((sheet.title, rows))
    except PDFMergeError:
        raise
    except Exception as exc:
        logger.warning("Reading xlsx rows failed: %s", exc)
        raise PDFMergeError("The Excel file is corrupted or unreadable.")
    finally:
        workbook.close()

    if not sheets:
        raise PDFMergeError("The spreadsheet contains no data.", status=422)

    # --- Render with reportlab ---
    font, font_bold = _register_pdf_font()
    page_size = landscape(A4)
    margin = 12 * mm
    avail_width = page_size[0] - 2 * margin
    styles = getSampleStyleSheet()
    title_style = styles["Heading2"]
    title_style.fontName = font_bold

    story = []
    for sheet_no, (title, rows) in enumerate(sheets):
        cols = len(rows[0])
        # Column widths from content length (capped), scaled to fit the page width
        font_size = 8
        widths = [
            max(6, min(40, max(len(r[c]) for r in rows))) * font_size * 0.55 + 8
            for c in range(cols)
        ]
        total = sum(widths)
        if total > avail_width:
            scale = avail_width / total
            widths = [w * scale for w in widths]
            font_size = max(5, font_size * scale)

        table = Table(rows, colWidths=widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("FONTNAME", (0, 0), (-1, 0), font_bold),
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
            ("LEADING", (0, 0), (-1, -1), font_size * 1.2),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFEFEF")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))

        if sheet_no:
            story.append(PageBreak())
        story.append(Paragraph(_xml_escape(title), title_style))
        story.append(Spacer(1, 4 * mm))
        story.append(table)

    output_path = os.path.join(work_dir, "output.pdf")
    try:
        doc = SimpleDocTemplate(
            output_path, pagesize=page_size,
            leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin,
            title=Path(input_path).stem, author="SnapPDF",
        )
        doc.build(story)
    except Exception as exc:   # reportlab LayoutError and friends
        logger.warning("reportlab build failed: %s", exc)
        raise PDFMergeError("The spreadsheet could not be laid out as a PDF.", status=500)

    return output_path, ".pdf", "application/pdf"


@csrf_exempt
@require_POST
def excel_to_pdf(request):
    """
    POST /excel-to-pdf/
    Form-data: file=<.xlsx or .xlsm>
    Returns: landscape PDF with one section per sheet, or JSON error.
    """
    return _run_single_file_tool(
        request, prefix="xlsx2pdf_", max_size=XLSX_MAX_FILE_SIZE,
        allowed_exts=XLSX_EXTENSIONS, magic_check=_is_zip_header, label="Excel",
        processor=_process_excel_to_pdf,
    )


# ----------------------------------------------------------------------
# Tool 12: Image to PDF (Pillow)
# ----------------------------------------------------------------------
def _process_image_to_pdf(request, input_path, work_dir):
    try:
        from PIL import Image, ImageOps, UnidentifiedImageError   # lazy import
    except ImportError:
        logger.error("Pillow not installed.")
        raise PDFMergeError("Server is missing image conversion tools.", status=500)

    output_path = os.path.join(work_dir, "output.pdf")
    try:
        with Image.open(input_path) as img:
            # Image.open is lazy: size is known before any pixel data is decoded
            width, height = img.size
            if width * height > IMAGE_MAX_PIXELS:
                raise PDFMergeError(
                    f"Image is {width}x{height}px; the maximum is {IMAGE_MAX_PIXELS:,} pixels.",
                    status=413,
                )

            # Physical size in the PDF follows the embedded DPI (defaults to 72)
            dpi_info = img.info.get("dpi") or (72, 72)
            try:
                dpi = float(dpi_info[0])
                if not 1 <= dpi <= 1200:
                    dpi = 72.0
            except (TypeError, ValueError, IndexError):
                dpi = 72.0

            try:
                img = ImageOps.exif_transpose(img) or img   # honour phone camera rotation
            except Exception:
                pass

            has_alpha = img.mode in ("RGBA", "LA", "PA") or (
                img.mode == "P" and "transparency" in img.info
            )
            if has_alpha:
                rgba = img.convert("RGBA")                  # flatten transparency onto white
                flat = Image.new("RGB", rgba.size, (255, 255, 255))
                flat.paste(rgba, mask=rgba.getchannel("A"))
                out = flat
            elif img.mode in ("RGB", "L", "CMYK"):
                out = img.copy()                            # PDF-safe modes as-is
            else:
                out = img.convert("RGB")                    # 1, P, I, F, YCbCr, ...

            out.save(output_path, "PDF", resolution=dpi)
            out.close()
    except PDFMergeError:
        raise
    except Image.DecompressionBombError:
        raise PDFMergeError("The image is too large to process safely.", status=413)
    except (UnidentifiedImageError, OSError, ValueError, SyntaxError) as exc:
        logger.warning("Image conversion failed: %s", exc)
        raise PDFMergeError("The image is corrupted or in an unsupported format.")

    return output_path, ".pdf", "application/pdf"


@csrf_exempt
@require_POST
def image_to_pdf(request):
    """
    POST /image-to-pdf/
    Form-data: file=<jpg|png|webp|bmp|gif|tiff>
    Returns: single-page PDF, or JSON error.
    """
    return _run_single_file_tool(
        request, prefix="img2pdf_", max_size=IMAGE_MAX_FILE_SIZE,
        allowed_exts=IMAGE_EXTENSIONS, magic_check=_is_image_header, label="image",
        processor=_process_image_to_pdf,
    )


# ----------------------------------------------------------------------
# Tool 13: Crop PDF (pypdf)
# ----------------------------------------------------------------------
_CROP_SIDES = ("left", "top", "right", "bottom")
_CROP_UNITS = {"pt": 1.0, "mm": 72 / 25.4, "cm": 72 / 2.54, "in": 72.0}


def _parse_crop_margins(post):
    """
    Read left/top/right/bottom margins (amount to REMOVE from each edge) plus an optional
    'unit' (pt, mm, cm, in; default pt). Returns margins in PDF points.
    """
    unit = (post.get("unit") or "pt").strip().lower()
    if unit not in _CROP_UNITS:
        raise PDFMergeError(f"'unit' must be one of: {', '.join(_CROP_UNITS)}.")
    factor = _CROP_UNITS[unit]

    missing = [side for side in _CROP_SIDES if post.get(side, "").strip() == ""]
    if missing:
        raise PDFMergeError(
            f"Missing crop value(s): {', '.join(missing)}. Provide left, top, right and bottom."
        )

    margins = {}
    for side in _CROP_SIDES:
        try:
            value = float(post[side])
        except (TypeError, ValueError):
            raise PDFMergeError(f"'{side}' must be a number.")
        if not math.isfinite(value) or value < 0:
            raise PDFMergeError(f"'{side}' must be a non-negative number.")
        margins[side] = value * factor

    if not any(margins.values()):
        raise PDFMergeError("At least one crop margin must be greater than zero.")
    return margins


def _margins_for_rotation(m, rotation):
    """
    Map margins given in the VIEWER's orientation onto the page's stored (unrotated)
    coordinates. Returns (left, bottom, right, top) to remove in page space.
    """
    l, t, r, b = m["left"], m["top"], m["right"], m["bottom"]
    if rotation == 90:
        return t, l, b, r
    if rotation == 180:
        return r, t, l, b
    if rotation == 270:
        return b, r, t, l
    return l, b, r, t


def _process_crop_pdf(request, input_path, work_dir):
    margins = _parse_crop_margins(request.POST)                    # fail fast on bad input

    raw_pages = (request.POST.get("pages") or "").strip()          # optional: "1,3,5-7"
    ranges = _parse_page_ranges(raw_pages) if raw_pages else None

    reader = _check_pdf_readable(input_path, CROP_MAX_PAGES)       # encrypted / corrupt / too big
    total_pages = len(reader.pages)
    indexes = _expand_page_ranges(ranges, total_pages) if ranges else range(total_pages)

    output_path = os.path.join(work_dir, "output.pdf")
    try:
        writer = PdfWriter(clone_from=reader)
    except (PyPdfError, ValueError, KeyError, OSError, RecursionError) as exc:
        logger.warning("Cloning PDF for crop failed: %s", exc)
        raise PDFMergeError("The PDF is corrupted or unreadable.")

    try:
        for index in indexes:
            page = writer.pages[index]
            try:
                rotation = int(page.rotation) % 360
            except Exception:
                rotation = 0

            ml, mb, mr, mt = _margins_for_rotation(margins, rotation)
            box = page.cropbox                                     # the area viewers display
            left = float(box.left) + ml
            bottom = float(box.bottom) + mb
            right = float(box.right) - mr
            top = float(box.top) - mt

            if right - left < 1 or top - bottom < 1:
                raise PDFMergeError(
                    f"Crop margins remove all of page {index + 1} "
                    f"({float(box.width):.0f} x {float(box.height):.0f} pt)."
                )

            new_box = RectangleObject((left, bottom, right, top))
            page.mediabox = new_box
            page.cropbox = new_box
            if "/TrimBox" in page:
                page.trimbox = new_box
            if "/BleedBox" in page:
                page.bleedbox = new_box
            if "/ArtBox" in page:
                page.artbox = new_box

        with open(output_path, "wb") as out:
            writer.write(out)
    except PDFMergeError:
        raise
    except (PyPdfError, ValueError, KeyError, IndexError, OSError, RecursionError) as exc:
        logger.warning("Crop failed: %s", exc)
        raise PDFMergeError("Some pages are corrupted and cannot be cropped.")
    finally:
        writer.close()

    return output_path, "_cropped.pdf", "application/pdf"


@csrf_exempt
@require_POST
def crop_pdf(request):
    """
    POST /crop-pdf/
    Form-data: file=<pdf>
               left, top, right, bottom = amount to remove from each edge (required)
               unit  = pt | mm | cm | in   (optional, default pt)
               pages = "1,3,5-7"           (optional, default all pages)
    Returns: cropped PDF, or JSON error.
    """
    return _run_single_pdf_tool(
        request, prefix="crop_", max_size=CROP_MAX_FILE_SIZE, processor=_process_crop_pdf
    )

# ======================================================================
# Tool 14: Edit PDF (text boxes + whiteout from the visual editor)
# ======================================================================
import json

EDIT_MAX_FILE_SIZE = getattr(settings, "PDF_EDIT_MAX_FILE_SIZE", 100 * 1024 * 1024)
EDIT_MAX_PAGES = getattr(settings, "PDF_EDIT_MAX_PAGES", 1000)
EDIT_MAX_ELEMENTS = getattr(settings, "PDF_EDIT_MAX_ELEMENTS", 500)
EDIT_MAX_TEXT_CHARS = getattr(settings, "PDF_EDIT_MAX_TEXT_CHARS", 5000)
EDIT_MAX_JSON_BYTES = getattr(settings, "PDF_EDIT_MAX_JSON_BYTES", 2 * 1024 * 1024)
EDIT_FONT_PATH = getattr(settings, "PDF_FONT_PATH", None)
EDIT_MIN_FONT_SIZE, EDIT_MAX_FONT_SIZE = 4.0, 200.0
EDIT_LINE_HEIGHT = 1.2            # MUST match the editor's CSS line-height for .txt

_HEX_COLOR_RE = re.compile(r"^#?([0-9a-fA-F]{6})$")
_CONTROL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _import_fitz():
    """PyMuPDF is importable as 'pymupdf' (new) or 'fitz' (legacy)."""
    try:
        import pymupdf as fitz
        return fitz
    except ImportError:
        try:
            import fitz
            return fitz
        except ImportError:
            logger.error("PyMuPDF is not installed.")
            raise PDFMergeError("Server is missing PDF editing tools.", status=500)


def _edit_font():
    """(fontname, fontfile) for PyMuPDF: Unicode TTF from settings if present, else Helvetica."""
    if EDIT_FONT_PATH and os.path.exists(EDIT_FONT_PATH):
        return "SnapPDFSans", EDIT_FONT_PATH
    return "helv", None


def _hex_to_rgb(value, label):
    match = _HEX_COLOR_RE.match(str(value).strip())
    if not match:
        raise PDFMergeError(f"{label}: invalid color '{value}'. Use #RRGGBB.")
    h = match.group(1)
    return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))


def _edit_number(obj, key, label, *, minimum=None, maximum=None, required=True):
    value = obj.get(key)
    if value is None:
        if required:
            raise PDFMergeError(f"{label}: '{key}' is required.")
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise PDFMergeError(f"{label}: '{key}' must be a number.")
    value = float(value)
    if not math.isfinite(value):
        raise PDFMergeError(f"{label}: '{key}' must be a finite number.")
    if minimum is not None and value < minimum:
        raise PDFMergeError(f"{label}: '{key}' must be at least {minimum}.")
    if maximum is not None and value > maximum:
        raise PDFMergeError(f"{label}: '{key}' must be at most {maximum}.")
    return value


def _parse_edit_payload(raw):
    """
    Validate the editor's JSON:
      {"version": 1, "redact": false,
       "elements": [
         {"page": 1, "type": "whiteout", "x": 72, "y": 90, "w": 200, "h": 30,
          "pageWidth": 612, "pageHeight": 792},
         {"page": 1, "type": "text", "x": 72, "y": 140, "w": 200, "h": 17,
          "text": "Hello", "fontSize": 14, "color": "#000000", ...}]}
    Coordinates are PDF points in the DISPLAYED page space, origin top-left, y down.
    Returns (elements, redact).
    """
    if not raw or not raw.strip():
        raise PDFMergeError("'edits' JSON is required.")
    if len(raw) > EDIT_MAX_JSON_BYTES:
        raise PDFMergeError("'edits' payload is too large.", status=413)
    try:
        data = json.loads(raw)
    except ValueError:
        raise PDFMergeError("'edits' is not valid JSON.")
    if not isinstance(data, dict) or not isinstance(data.get("elements"), list):
        raise PDFMergeError("'edits' must be an object with an 'elements' list.")

    raw_elements = data["elements"]
    if not raw_elements:
        raise PDFMergeError("No edits to apply. Add text or a whiteout box first.")
    if len(raw_elements) > EDIT_MAX_ELEMENTS:
        raise PDFMergeError(f"Too many edits; the maximum is {EDIT_MAX_ELEMENTS}.", status=413)

    elements = []
    for index, el in enumerate(raw_elements, start=1):
        label = f"Element {index}"
        if not isinstance(el, dict):
            raise PDFMergeError(f"{label}: must be an object.")
        kind = str(el.get("type", "")).lower()
        if kind not in ("text", "whiteout"):
            raise PDFMergeError(f"{label}: 'type' must be 'text' or 'whiteout'.")

        item = {
            "type": kind,
            "page": int(_edit_number(el, "page", label, minimum=1, maximum=EDIT_MAX_PAGES)),
            "x": _edit_number(el, "x", label, minimum=-100_000, maximum=100_000),
            "y": _edit_number(el, "y", label, minimum=-100_000, maximum=100_000),
            "w": _edit_number(el, "w", label, minimum=0, maximum=100_000),
            "h": _edit_number(el, "h", label, minimum=0, maximum=100_000),
            "page_w": _edit_number(el, "pageWidth", label, minimum=1, maximum=100_000, required=False),
            "page_h": _edit_number(el, "pageHeight", label, minimum=1, maximum=100_000, required=False),
        }

        if kind == "text":
            text = el.get("text", "")
            if not isinstance(text, str):
                raise PDFMergeError(f"{label}: 'text' must be a string.")
            text = _CONTROL_CHARS_RE.sub("", text.replace("\r\n", "\n").replace("\r", "\n"))
            if not text.strip():
                raise PDFMergeError(f"{label}: 'text' is empty.")
            if len(text) > EDIT_MAX_TEXT_CHARS:
                raise PDFMergeError(
                    f"{label}: text exceeds {EDIT_MAX_TEXT_CHARS} characters.", status=413
                )
            font_size = _edit_number(
                el, "fontSize", label,
                minimum=EDIT_MIN_FONT_SIZE, maximum=EDIT_MAX_FONT_SIZE, required=False,
            )
            align = str(el.get("align", "left")).lower()
            if align not in ("left", "center", "right"):
                raise PDFMergeError(f"{label}: 'align' must be left, center or right.")
            item.update(
                text=text,
                font_size=font_size if font_size is not None else 14.0,
                color=_hex_to_rgb(el.get("color", "#000000"), label),
                bold=bool(el.get("bold", False)),
                align=align,
            )
        else:  # whiteout
            if item["w"] <= 0 or item["h"] <= 0:
                raise PDFMergeError(
                    f"{label}: a whiteout box needs a positive width and height."
                )
            item["fill"] = _hex_to_rgb(el.get("fill", "#ffffff"), label)

        elements.append(item)

    return elements, bool(data.get("redact", False))


# ----------------------------------------------------------------------
# Geometry: editor coordinates -> PyMuPDF rectangles
# ----------------------------------------------------------------------
def _element_rect(fitz, el, page_rect, label):
    """
    Convert an element's editor coordinates (PDF points, DISPLAYED page, origin top-left,
    y down) into a fitz.Rect on the displayed page. pdf.js viewports and page.rect share
    this convention, so this is normally a 1:1 copy. If the client measured a different
    page size (pageWidth/pageHeight), rescale proportionally. Finally clip to the page.
    """
    sx = page_rect.width / el["page_w"] if el["page_w"] else 1.0
    sy = page_rect.height / el["page_h"] if el["page_h"] else 1.0
    box = fitz.Rect(
        page_rect.x0 + el["x"] * sx,
        page_rect.y0 + el["y"] * sy,
        page_rect.x0 + (el["x"] + el["w"]) * sx,
        page_rect.y0 + (el["y"] + el["h"]) * sy,
    )
    clipped = box & page_rect
    if clipped.is_empty or clipped.width < 0.5 or clipped.height < 0.5:
        raise PDFMergeError(f"{label} lies outside page {el['page']}.")
    return clipped


def _draw_whiteout(fitz, page, box, fill):
    """Paint an opaque rectangle over existing content (content underneath is hidden, not removed)."""
    shape = page.new_shape()
    # Drawing happens in the page's UNROTATED space: derotate the displayed-space box.
    shape.draw_rect(box * page.derotation_matrix)
    shape.finish(color=None, fill=fill, width=0)
    shape.commit()   # overlay=True by default: drawn on top of the page content


def _redact_boxes(fitz, page, boxes):
    """True redaction: physically remove text/images under each box, then fill it."""
    derot = page.derotation_matrix
    for box, fill in boxes:
        page.add_redact_annot(box * derot, fill=fill)
    page.apply_redactions(images=getattr(fitz, "PDF_REDACT_IMAGE_PIXELS", 2))


def _textbox(page, rect, text, font_size, kwargs):
    """insert_textbox with the editor's line height; older PyMuPDF lacks 'lineheight'."""
    try:
        return page.insert_textbox(
            rect, text, fontsize=font_size, lineheight=EDIT_LINE_HEIGHT, **kwargs
        )
    except TypeError:
        return page.insert_textbox(rect, text, fontsize=font_size, **kwargs)


def _insert_text(fitz, page, box, el, fontname, fontfile, label):
    """
    Write a text box at 'box' (displayed-page coordinates). If the text does not fit
    (browser and PDF font metrics differ slightly), first grow the box downward within
    the page, then shrink the font, so the user's text is never silently dropped.
    """
    align = {
        "left": fitz.TEXT_ALIGN_LEFT,
        "center": fitz.TEXT_ALIGN_CENTER,
        "right": fitz.TEXT_ALIGN_RIGHT,
    }[el["align"]]

    kwargs = dict(
        fontname=fontname,
        fontfile=fontfile,
        color=el["color"],
        align=align,
        rotate=page.rotation,   # keeps text upright on /Rotate 90/180/270 pages
    )
    if el["bold"]:
        if fontfile is None:
            kwargs["fontname"] = "hebo"          # Helvetica-Bold
        else:
            kwargs["render_mode"] = 2            # fill + stroke = synthetic bold for TTF
            kwargs["border_width"] = max(0.2, el["font_size"] * 0.025)

    page_rect = page.rect
    derot = page.derotation_matrix
    font_size = el["font_size"]

    for _ in range(12):
        rc = _textbox(page, box * derot, el["text"], font_size, kwargs)
        if rc >= 0:
            return                                # written successfully
        # rc < 0 is the missing height: try growing the box downward inside the page
        grown = fitz.Rect(box.x0, box.y0, box.x1, min(page_rect.y1, box.y1 - rc + 1))
        if grown.height > box.height + 0.5:
            box = grown
            continue
        font_size *= 0.9                          # cannot grow: shrink the font
        if font_size < EDIT_MIN_FONT_SIZE:
            break

    raise PDFMergeError(
        f"{label}: the text does not fit on page {el['page']}. "
        "Enlarge the box or shorten the text."
    )


# ----------------------------------------------------------------------
# Processor
# ----------------------------------------------------------------------
def _process_edit_pdf(request, input_path, work_dir):
    elements, redact = _parse_edit_payload(request.POST.get("edits"))   # fail fast on bad JSON
    _check_pdf_readable(input_path, EDIT_MAX_PAGES)                      # encrypted / corrupt / too big

    fitz = _import_fitz()
    fontname, fontfile = _edit_font()

    try:
        doc = fitz.open(input_path)
    except Exception as exc:
        logger.warning("PyMuPDF could not open PDF: %s", exc)
        raise PDFMergeError("The PDF is corrupted or unreadable.")

    try:
        if doc.needs_pass:
            raise PDFMergeError("The PDF is password protected. Unlock it first.")
        page_count = doc.page_count

        # Group elements by page and validate page numbers against the real document
        by_page = {}
        for index, el in enumerate(elements, start=1):
            if el["page"] > page_count:
                raise PDFMergeError(
                    f"Element {index} targets page {el['page']}, "
                    f"but the PDF has {page_count} page(s)."
                )
            el["label"] = f"Element {index}"
            by_page.setdefault(el["page"], []).append(el)

        for page_no, items in sorted(by_page.items()):
            page = doc[page_no - 1]
            page_rect = page.rect   # displayed page, rotation applied, origin (0, 0) top-left
            whiteouts = [e for e in items if e["type"] == "whiteout"]
            texts = [e for e in items if e["type"] == "text"]

            # 1. Whiteouts first, so text placed over them stays visible (matches editor z-order)
            if whiteouts:
                boxes = [
                    (_element_rect(fitz, e, page_rect, e["label"]), e["fill"]) for e in whiteouts
                ]
                if redact:
                    _redact_boxes(fitz, page, boxes)
                else:
                    for box, fill in boxes:
                        _draw_whiteout(fitz, page, box, fill)

            # 2. Text boxes
            for e in texts:
                box = _element_rect(fitz, e, page_rect, e["label"])
                _insert_text(fitz, page, box, e, fontname, fontfile, e["label"])

        output_path = os.path.join(work_dir, "output.pdf")
        doc.save(output_path, garbage=3, deflate=True)

    except PDFMergeError:
        raise
    except Exception as exc:   # MuPDF raises RuntimeError/ValueError for damaged structures
        logger.warning("Edit PDF failed: %s", exc)
        raise PDFMergeError("The PDF could not be edited. It may be corrupted.")
    finally:
        doc.close()

    return output_path, "_edited.pdf", "application/pdf"


# ----------------------------------------------------------------------
# View
# ----------------------------------------------------------------------
@csrf_exempt
@require_POST
def edit_pdf(request):
    """
    POST /api/edit-pdf/
    Form-data: file=<pdf>
               edits=<JSON string, see _parse_edit_payload>
    Returns: edited PDF as an attachment, or JSON {"error": "..."}.
    The temp directory is deleted after the response by SelfCleaningFileResponse.
    """
    return _run_single_pdf_tool(
        request, prefix="edit_", max_size=EDIT_MAX_FILE_SIZE, processor=_process_edit_pdf
    )


